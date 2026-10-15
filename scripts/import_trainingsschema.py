#!/usr/bin/env python3
"""Zet het trainingsschema van het bestuur (Excel) om in data/trainingen.json.

Gebruik:  python3 scripts/import_trainingsschema.py <trainingsschema.xlsx>
Vereist:  pip install openpyxl

Verwacht de indeling van het bestuur: op blad 1 per dag blokken met de tijd in
kolom B (De Tas, veld 1 t/m 4 in C t/m F) en kolom H (De Bocht, veld 4 t/m 7 in
I t/m L), met de dagnaam in kolom A. Een cel kan een team zijn, een team met een
eigen tijd ("D1 (1) 19.00-20.45"), "vanaf 19.00 3 velden", "X" of "Wedstrijd ...".
De veldnummers staan in de kopregel (bijvoorbeeld "Veld 4 (LV)") en worden
overgenomen, zodat ze overeenkomen met het wedstrijdschema: De Tas 1 t/m 3 voor
wedstrijden (trainingen gebruiken soms ook 4), De Bocht 4 t/m 7.
"""
import json, os, re, sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DAGEN = {'maandag': 'ma', 'dinsdag': 'di', 'woensdag': 'wo', 'donderdag': 'do', 'vrijdag': 'vr'}
TIJD = re.compile(r'(\d{1,2})[.:](\d{2})\s*-\s*(\d{1,2})[.:](\d{2})')


def hhmm(u, m):
    return '%02d:%02d' % (int(u), int(m))


def blok_tijd(tekst):
    m = TIJD.search(str(tekst or ''))
    return (hhmm(m.group(1), m.group(2)), hhmm(m.group(3), m.group(4))) if m else None


def cel(tekst):
    t = str(tekst).strip()
    if t.lower().startswith('wedstrijd'):
        return {'team': None, 'wedstrijd': True, 'opm': t}
    if t.upper() == 'X':
        return {'team': None, 'opm': 'X (niet beschikbaar)'}
    if t.lower().startswith('vanaf'):
        return {'team': None, 'opm': t}
    m = re.match(r'^(.*?)\s+(\d{1,2}[.:]\d{2}\s*-\s*\d{1,2}[.:]\d{2})$', t)
    if m:
        van, tot = blok_tijd(m.group(2))
        return {'team': m.group(1).strip(), 'van': van, 'tot': tot}
    return {'team': t}


def main(pad):
    ws = openpyxl.load_workbook(pad, data_only=True).worksheets[0]
    versie = str(ws['A1'].value or '').replace('Versie', '').strip()
    # kopregel: kolom -> veldnummer, per zaal
    kop = {}
    for rij in ws.iter_rows(min_row=1, max_row=6):
        for c in rij:
            m = re.match(r'^\s*Veld\s+(\d+)\s*\((\w+)\)', str(c.value or ''))
            if m:
                kop[c.column] = (m.group(1), m.group(2))
    zalen = {}                  # kolom -> 'Tas' / 'Bocht'
    for c in ws[3]:
        if str(c.value or '').lower().startswith('sporthal'):
            zalen[c.column] = 'Bocht' if 'bocht' in c.value.lower() else 'Tas'
    kolommen = {'Tas': sorted(c for c in kop if c < 7), 'Bocht': sorted(c for c in kop if c >= 7)}
    tijdkolom = {'Tas': 2, 'Bocht': 8}

    dagen, dag, opmerkingen, veldnamen = {}, None, [], {'Tas': {}, 'Bocht': {}}
    for hal in kolommen:
        for k in kolommen[hal]:
            veldnamen[hal][kop[k][0]] = {'RV': 'rechtsvoor', 'LV': 'linksvoor', 'LA': 'linksachter', 'RA': 'rechtsachter'}.get(kop[k][1], kop[k][1])
    for r in range(5, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or '').strip().lower()
        if a in DAGEN:
            dag = DAGEN[a]
        if dag is None:
            continue
        for hal in ('Tas', 'Bocht'):
            tijd = blok_tijd(ws.cell(r, tijdkolom[hal]).value)
            if not tijd:
                continue
            velden = {}
            for k in kolommen[hal]:
                v = ws.cell(r, k).value
                if v not in (None, ''):
                    velden[kop[k][0]] = cel(v)
            if velden:
                dagen.setdefault(dag, []).append({'hal': hal, 'van': tijd[0], 'tot': tijd[1], 'velden': velden})
        tekst = str(ws.cell(r, 9).value or '')
        if tekst.lower().startswith('bij wedstrijddagen'):
            opmerkingen.append(tekst.strip() + '.')

    uit = {
        'versie': versie,
        'toelichting': 'Trainingsschema van De Fluit, overgenomen uit het Excel van het bestuur. Cijfer tussen haakjes = eerste of tweede training van het team. De Tas heeft bij trainingen veld 1 t/m 4, bij wedstrijden 1 t/m 3; De Bocht heeft veld 4 t/m 7, net als in het wedstrijdschema.',
        'opmerkingen': opmerkingen,
        'wedstrijdDuurMinuten': 120,
        'velden': veldnamen,
        'dagen': dagen,
    }
    doel = os.path.join(ROOT, 'data/trainingen.json')
    with open(doel, 'w', encoding='utf-8') as f:
        json.dump(uit, f, ensure_ascii=False, indent=1)
        f.write('\n')
    print('versie %s: %d dagen, %d blokken, %d cellen' % (versie, len(dagen), sum(len(b) for b in dagen.values()), sum(len(x['velden']) for b in dagen.values() for x in b)))


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
